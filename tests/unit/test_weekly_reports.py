from datetime import datetime, timezone
from io import BytesIO

from tp_codex import weekly_reports
from tp_codex.workbooks import WorkbookRow, WorkbookSheet, _build_xlsx
from openpyxl import load_workbook


def test_resolve_week_context_defaults_to_previous_iso_week_in_shanghai_time():
    context = weekly_reports.resolve_week_context(None, "2026-06-08T01:00:00+00:00")

    assert context.report_year == 2026
    assert context.week_number == 23
    assert context.dataset_week == "2026-W23"
    assert context.week_label == "Week23"
    assert context.start_date.isoformat() == "2026-06-01"
    assert context.end_date.isoformat() == "2026-06-07"
    assert context.output_filename == "质量周报-Week23(2026.6.1-2026.6.7).xlsx"


def test_build_weekly_created_where_uses_datetime_parse_window():
    context = weekly_reports.resolve_week_context("2026-W23", datetime(2026, 6, 6, tzinfo=timezone.utc).isoformat())

    assert (
        weekly_reports.build_weekly_created_where(context)
        == 'CreateDate >= DateTime.Parse("2026-06-01") and CreateDate < DateTime.Parse("2026-06-08")'
    )


def test_build_di_summary_weights_detected_and_remaining_effective_bugs():
    records = [
        {"severity": "Blocking", "status_raw": "New", "status_group": "triage"},
        {"severity": "Critical", "status_raw": "Verified", "status_group": "closed"},
        {"severity": "Major", "status_raw": "Invalid", "status_group": "closed"},
        {"severity": "Normal", "status_raw": "In Progress", "status_group": "in_progress"},
    ]

    summary = weekly_reports.build_di_summary(
        records,
        {"di_weights": {"Blocking": 10, "Critical": 3, "Major": 1, "Normal": 0.1}},
    )

    assert summary == {
        "detected_di": 13.1,
        "remaining_di": 10.1,
        "closed_rate": 3 / 13.1,
        "by_severity": {
            "Blocking": {"detected_count": 1, "remaining_count": 1, "detected_di": 10, "remaining_di": 10},
            "Critical": {"detected_count": 1, "remaining_count": 0, "detected_di": 3, "remaining_di": 0},
            "Normal": {"detected_count": 1, "remaining_count": 1, "detected_di": 0.1, "remaining_di": 0.1},
        },
    }


def test_append_weekly_sheet_rewrites_existing_target_week_sheet():
    template = _build_xlsx(
        [
            WorkbookSheet(
                name="固件质量数据概览-Week23",
                rows=_weekly_template_rows("Week23新增61个Bug", 61),
            ),
            WorkbookSheet(name="辅助sheet", rows=[WorkbookRow(["辅助"], kind="title")]),
        ],
        "2026-06-05T00:00:00+00:00",
    )
    context = weekly_reports.resolve_week_context("Week23", "2026-06-08T03:58:27.648398+00:00")
    records = [
        {
            "bug_id": 1,
            "name": "NG3 blocking",
            "team": "ESW China NG3 Driver",
            "products": [],
            "severity": "Blocking",
            "status_raw": "New",
            "status_group": "triage",
            "created_at": "2026-06-02 10:00:00",
            "created_week": "2026-W23",
            "is_customer_feedback": False,
            "is_open": True,
            "is_closed": False,
            "url": "https://example.com/1",
        },
        {
            "bug_id": 2,
            "name": "NG3 major",
            "team": "ESW UI Team",
            "products": [],
            "severity": "Major",
            "status_raw": "In Progress",
            "status_group": "in_progress",
            "created_at": "2026-06-03 10:00:00",
            "created_week": "2026-W23",
            "is_customer_feedback": True,
            "is_open": True,
            "is_closed": False,
            "url": "https://example.com/2",
        },
    ]

    workbook_bytes = weekly_reports.append_weekly_sheet(
        template,
        records,
        context,
        weekly_records=records,
    )

    workbook = load_workbook(BytesIO(workbook_bytes))
    assert workbook.sheetnames == ["固件质量数据概览-Week23"]
    sheet = workbook["固件质量数据概览-Week23"]
    assert "Week23新增2个Bug" in str(sheet["A3"].value)
    assert sheet["B9"].value == 2


def test_append_weekly_sheet_includes_wui_scope_and_rolls_it_into_framework_rows():
    template = _build_xlsx(
        [
            WorkbookSheet(
                name="固件质量数据概览-Week23",
                rows=_weekly_template_rows("Week23新增61个Bug", 61),
            )
        ],
        "2026-06-05T00:00:00+00:00",
    )
    context = weekly_reports.resolve_week_context("Week23", "2026-06-08T03:58:27.648398+00:00")
    records = [
        {
            "bug_id": 1,
            "name": "Driver blocking",
            "team": "ESW China NG3 Driver",
            "products": [],
            "severity": "Blocking",
            "status_raw": "New",
            "status_group": "triage",
            "created_at": "2026-06-02 10:00:00",
            "created_week": "2026-W23",
            "is_customer_feedback": False,
            "is_open": True,
            "is_closed": False,
            "url": "https://example.com/1",
        },
        {
            "bug_id": 2,
            "name": "WUI critical",
            "team": "ESW WUI",
            "products": [],
            "severity": "Critical",
            "status_raw": "In Progress",
            "status_group": "in_progress",
            "created_at": "2026-06-03 10:00:00",
            "created_week": "2026-W23",
            "is_customer_feedback": False,
            "is_open": True,
            "is_closed": False,
            "url": "https://example.com/2",
        },
        {
            "bug_id": 3,
            "name": "WUI customer feedback",
            "team": "ESW WUI",
            "products": [],
            "severity": "Major",
            "status_raw": "Verified",
            "status_group": "closed",
            "created_at": "2026-06-04 10:00:00",
            "created_week": "2026-W23",
            "is_customer_feedback": True,
            "is_open": False,
            "is_closed": True,
            "url": "https://example.com/3",
        },
        {
            "bug_id": 4,
            "name": "Historic WUI verified stock bug",
            "team": "ESW WUI",
            "products": [],
            "severity": "Major",
            "status_raw": "Verified",
            "status_group": "closed",
            "created_at": "2025-06-04 10:00:00",
            "created_week": "2025-W23",
            "last_status_change_at": "2026-06-04 10:00:00",
            "is_customer_feedback": False,
            "is_open": False,
            "is_closed": True,
            "url": "https://example.com/4",
        },
    ]
    weekly_records = records[:3]

    workbook_bytes = weekly_reports.append_weekly_sheet(
        template,
        records,
        context,
        weekly_records=weekly_records,
    )

    workbook = load_workbook(BytesIO(workbook_bytes))
    sheet = workbook["固件质量数据概览-Week23"]
    assert "框架2个" in str(sheet["A3"].value)
    assert sheet["B9"].value == 3
    assert sheet["A21"].value == "框架"
    assert sheet["B21"].value == 1
    assert sheet["A34"].value == "UI"
    assert sheet["B34"].value == 0
    assert sheet["D34"].value == 0
    assert sheet["E34"].value == 0
    assert sheet["A40"].value == "框架"
    assert sheet["B40"].value == 0
    assert sheet["C40"].value == 1


def test_effective_bug_table_uses_allowed_states_and_excludes_customer_feedback_duplicate_invalid():
    template = _build_xlsx(
        [
            WorkbookSheet(
                name="固件质量数据概览-Week23",
                rows=_weekly_template_rows("说明", 0),
            )
        ],
        "2026-06-05T00:00:00+00:00",
    )
    context = weekly_reports.resolve_week_context("Week23", "2026-06-08T03:58:27.648398+00:00")
    base_record = {
        "team": "ESW China NG3 Driver",
        "products": [],
        "severity": "Critical",
        "status_group": "closed",
        "created_at": "2026-02-01 10:00:00",
        "created_week": "2026-W05",
        "is_customer_feedback": False,
        "is_open": False,
        "is_closed": True,
        "url": "https://example.com/1",
    }
    allowed_states = [
        "New",
        "In Progress",
        "In Review",
        "In Testing",
        "Reproduce",
        "Verified",
        "Wont fix",
        "Fixed",
        "Expired",
        "Later",
        "Planned",
    ]
    records = [
        {**base_record, "bug_id": index, "name": f"effective {state}", "status_raw": state}
        for index, state in enumerate(allowed_states, start=1)
    ]
    records.extend(
        [
            {
                **base_record,
                "bug_id": 99,
                "name": "race product still belongs to NG3 platform",
                "status_raw": "Verified",
                "products": ["Suunto Race 3", "Suunto Race 3 S"],
            },
            {**base_record, "bug_id": 100, "name": "Customer feedback should be excluded", "status_raw": "Verified", "is_customer_feedback": True},
            {**base_record, "bug_id": 101, "name": "duplicate should be excluded", "status_raw": "Duplicate"},
            {**base_record, "bug_id": 102, "name": "invalid should be excluded", "status_raw": "Invalid"},
            {**base_record, "bug_id": 103, "name": "unknown should be excluded", "status_raw": "Rejected"},
        ]
    )

    workbook_bytes = weekly_reports.append_weekly_sheet(template, records, context, weekly_records=[])

    workbook = load_workbook(BytesIO(workbook_bytes))
    sheet = workbook["固件质量数据概览-Week23"]
    assert sheet["B18"].value == len(allowed_states) + 1
    assert sheet["D18"].value == 1
    assert sheet["E18"].value == 3
    assert sheet["B15"].value == len(allowed_states) + 1
    assert sheet["D15"].value == 1
    assert sheet["E15"].value == 3
    assert "有效Bug总数12个" in sheet["A12"].value


def _weekly_template_rows(summary_text: str, total: int) -> list[WorkbookRow]:
    cells = {
        1: ["固件质量报告2026 Week23"],
        2: ["2026年NG3固件每周新增Bug"],
        3: [summary_text],
        4: ["等级", "总数", "New", "处理中", "已解决", "已验证", "异常闭环"],
        5: ["Blocking", 0, 0, 0, 0, 0, 0],
        6: ["Critical", 0, 0, 0, 0, 0, 0],
        7: ["Major", 0, 0, 0, 0, 0, 0],
        8: ["Normal", 0, 0, 0, 0, 0, 0],
        9: ["总计", total, 0, 0, 0, 0, 0],
        10: ["各状态问题占比", "/", 0, 0, 0, 0, 0],
        11: ["NG3项目2026年固件有效bug检出&修复情况"],
        12: ["说明"],
        13: ["等级", "总数", "待解决", "待验证", "非常规闭环", "关闭率", "验证率"],
        14: ["Blocking", 0, 0, 0, 0, 0, 0],
        15: ["Critical", 0, 0, 0, 0, 0, 0],
        16: ["Major", 0, 0, 0, 0, 0, 0],
        17: ["Normal", 0, 0, 0, 0, 0, 0],
        18: ["总计", 0, 0, 0, 0, 0, 0],
        19: ["工作组", "总数", "待解决", "待验证", "非常规闭环", "关闭率", "验证率"],
        20: ["驱动", 0, 0, 0, 0, 0, 0],
        21: ["框架", 0, 0, 0, 0, 0, 0],
        22: ["UI", 0, 0, 0, 0, 0, 0],
        23: ["2026年NG3固件售后问题"],
        24: ["说明"],
        25: ["等级", "总数", "已闭环", "待关闭", "待验证", "非常规闭环", "关闭率"],
        26: ["Blocking", 0, 0, 0, 0, 0, 0],
        27: ["Critical", 0, 0, 0, 0, 0, 0],
        28: ["Major", 0, 0, 0, 0, 0, 0],
        29: ["Normal", 0, 0, 0, 0, 0, 0],
        30: ["总计", 0, 0, 0, 0, 0, 0],
        31: ["工作组", "Blocking总数", "Blocking超期", "Critical总数", "Critical待解决", "备注"],
        32: ["驱动", 0, 0, 0, 0, ""],
        33: ["框架", 0, 0, 0, 0, ""],
        34: ["UI", 0, 0, 0, 0, ""],
        35: ["总计", 0, 0, 0, 0, "/"],
        36: ["NG3存量Bug消减情况"],
        37: ["说明"],
        38: ["工作组", "Bug存量", "2026年关闭数", "待研发处理", "待复现", "待验证", "消减率"],
        39: ["驱动", 0, 0, 0, 0, 0, 0],
        40: ["框架", 0, 0, 0, 0, 0, 0],
        41: ["UI", 0, 0, 0, 0, 0, 0],
        42: ["总计", 0, 0, 0, 0, 0, 0],
    }
    return [WorkbookRow(cells.get(index, []), kind="body") for index in range(1, 43)]
